"""Independent reconciliation of normalized prices to raw cells and annual December table."""
import json, hashlib
from pathlib import Path
import openpyxl
root=Path(__file__).resolve().parents[1]
h=json.loads((root/'data/official-history.json').read_text())
books={s['id']:openpyxl.load_workbook(root/s['file'],read_only=False,data_only=True) for s in h['sources'] if s['file'].endswith('.xlsx')}
for s in h['sources']:
 assert hashlib.sha256((root/s['file']).read_bytes()).hexdigest()==s['sha256']
count=0
for group,scale in [(h['producer']['monthly'],1000),(h['producer']['annual'],1000),(h['retail']['weekly'],1),(h['producer']['regional'],1000)]:
 for row in group:
  original=books[row['sourceId']][row['sheet']][row['cell']].value
  assert original==row['rawValue']
  assert abs(original/scale-row['value'])<0.00001
  count+=1
# Cross-table check: December monthly values against independent end-of-year table.
s=books[h['sources'][0]['id']]['1.2']
r=next(r for r in range(1,s.max_row+1) if s.cell(r,2).value=='01.13.43.110')
checks=0
for c in range(3,s.max_column+1):
 year=str(s.cell(4,c).value)[:4]
 row=next((r for r in h['producer']['monthly'] if r['period']==year+'-12'),None)
 if row:
  assert abs(s.cell(r,c).value/1000-row['value'])<.001
  checks+=1
assert checks==5
print(f'PASS: {count} price cells reconciled; {checks} December cross-table checks; all source SHA-256 hashes match.')
